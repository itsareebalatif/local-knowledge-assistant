from __future__ import annotations

import hashlib
import io

import pytest
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool


@pytest.fixture(autouse=True)
def _disable_langfuse_tracing(monkeypatch):
    from app.config import get_settings
    from app.observability import langfuse_client

    monkeypatch.setattr(get_settings(), "langfuse_secret_key", None, raising=False)
    monkeypatch.setattr(get_settings(), "langfuse_public_key", None, raising=False)
    langfuse_client.get_langfuse.cache_clear()
    yield
    langfuse_client.get_langfuse.cache_clear()


@pytest.fixture
def db_session():
    """A fresh, isolated in-memory SQLite DB per test, with FK enforcement
    on (SQLite ignores foreign keys unless told otherwise, same as the real
    engine in app/db/base.py). Importing app.models registers every table on
    Base.metadata before create_all runs."""
    from app import models  # noqa: F401 - registers tables on Base.metadata
    from app.db.base import Base
    from app.db.fts import create_fts_index

    engine = create_engine(
        "sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )

    @event.listens_for(engine, "connect")
    def _enable_fk(dbapi_connection, _connection_record):
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    Base.metadata.create_all(engine)
    create_fts_index(engine)
    session = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture
def user(db_session):
    from app.models import User, UserAuth

    u = User(email="mohsin@example.com", full_name="Mohsin", role="ADMIN")
    db_session.add(u)
    db_session.flush()
    db_session.add(UserAuth(user_id=u.user_id, password_hash="hashed"))
    db_session.commit()
    return u


@pytest.fixture
def registered_user(db_session):
    from app.auth.auth_service import register_user

    return register_user(
        db_session, email="registered@example.com", password="correct horse battery staple", full_name="Registered User"
    )


@pytest.fixture
def auth_headers(registered_user) -> dict:
    return {"Authorization": f"Bearer {registered_user.access_token}"}


@pytest.fixture
def sample_txt_bytes() -> bytes:
    text = (
        "This is the first paragraph of a plain text file. It has a couple "
        "of sentences to make sure paragraph joining works.\n\n"
        "This is the second paragraph, completely separate from the first."
    )
    return text.encode("utf-8")


@pytest.fixture
def sample_markdown_bytes() -> bytes:
    text = """# Personal Knowledge Engine

Intro paragraph explaining the project in a sentence or two.

## Architecture

The system uses a vector store and a graph store together.

### Vector Store

Chunks are embedded with nomic-embed-text and stored in ChromaDB.

## Deployment

| Component | Tool |
| --- | --- |
| Vector DB | ChromaDB |
| Graph DB | NetworkX |

```python
def hello():
    return "world"
```
"""
    return text.encode("utf-8")


@pytest.fixture
def sample_html_bytes() -> bytes:
    html = """<!doctype html>
<html><head><title>Test Doc</title>
<style>body { color: red; }</style>
</head>
<body>
<nav>Site Nav Should Be Stripped</nav>
<header>Header Should Be Stripped</header>
<h1>Main Title</h1>
<p>Introductory paragraph with some content about the project.</p>
<h2>Sub Section</h2>
<p>Sub-section paragraph text goes here for testing purposes.</p>
<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>
<footer>Footer Should Be Stripped</footer>
<script>console.log('should be stripped')</script>
</body></html>
"""
    return html.encode("utf-8")


@pytest.fixture
def sample_docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Main Title", level=1)
    document.add_paragraph("Introductory paragraph with some content about the project.")
    document.add_heading("Sub Section", level=2)
    document.add_paragraph("Sub-section paragraph text goes here for testing purposes.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "2"
    document.add_paragraph("A closing paragraph after the table.")

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    content_width = pdf.w - pdf.l_margin - pdf.r_margin
    pdf.set_font("Helvetica", "B", 20)
    pdf.multi_cell(content_width, 12, "Main Title")
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(content_width, 8, "Introductory paragraph with some content about the project, long enough to wrap.")
    pdf.set_font("Helvetica", "B", 15)
    pdf.multi_cell(content_width, 10, "Sub Section")
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(content_width, 8, "Sub-section paragraph text goes here for testing purposes.")
    return bytes(pdf.output())


@pytest.fixture
def oversized_paragraph_text() -> str:
    # ~900 words with no punctuation, to force the hard-wrap fallback path.
    return " ".join(f"word{i}" for i in range(900))


class FakeEmbedder:
    _DIMS = 64

    def __init__(self):
        self.calls: list[list[str]] = []
        self.model = "fake-embedder"

    async def embed(self, texts: list[str], input_type: str = "search_document") -> list[list[float]]:
        self.calls.append(list(texts))
        return [self._vectorize(t) for t in texts]

    @classmethod
    def _vectorize(cls, text: str) -> list[float]:
        vec = [0.0] * cls._DIMS
        words = [w for w in text.lower().split() if w]
        for word in words:
            idx = int(hashlib.md5(word.encode()).hexdigest(), 16) % cls._DIMS
            vec[idx] += 1.0
        if not words:
            vec[0] = 1.0  # avoid an all-zero vector for empty input
        return vec


class FakeVectorStore:
    def __init__(self):
        self.records: dict[str, dict] = {}

    def add(self, ids, embeddings, documents, metadatas):
        for i, vec, doc, meta in zip(ids, embeddings, documents, metadatas):
            self.records[i] = {"embedding": vec, "document": doc, "metadata": meta}

    def query(self, embedding, top_k=5, where=None):
        items = list(self.records.items())
        if where:
            items = [(vec_id, rec) for vec_id, rec in items if all(rec["metadata"].get(k) == v for k, v in where.items())]
        return [
            {"id": vec_id, "document": rec["document"], "distance": float(rank), "metadata": rec["metadata"]}
            for rank, (vec_id, rec) in enumerate(items[:top_k])
        ]

    def delete(self, ids):
        for vec_id in ids:
            self.records.pop(vec_id, None)


class FakeLLM:
    """Deterministic stand-in for OllamaLLM/OpenAICompatibleLLM — no network, no live
    model server needed. Yields a fixed list of pieces and records every
    (system_prompt, user_prompt) call for assertions."""

    def __init__(self, pieces: list[str] | None = None):
        self.pieces = pieces if pieces is not None else ["This ", "is ", "a ", "fake ", "answer."]
        self.calls: list[tuple[str, str]] = []
        self.model = "fake-llm"

    async def generate_stream(self, system_prompt: str, user_prompt: str):
        self.calls.append((system_prompt, user_prompt))
        for piece in self.pieces:
            yield piece


class FakeReranker:
    def __init__(self):
        self.calls: list[tuple[str, int]] = []  # (query, num_candidates_in)

    def rerank(self, query: str, candidates: list, top_k: int) -> list:
        self.calls.append((query, len(candidates)))
        return list(reversed(candidates))[:top_k]


@pytest.fixture
def fake_reranker() -> FakeReranker:
    return FakeReranker()


@pytest.fixture
def fake_llm() -> FakeLLM:
    return FakeLLM()


@pytest.fixture
def fake_embedder() -> FakeEmbedder:
    return FakeEmbedder()


class FakeEmailSender:
    def __init__(self):
        self.calls: list[tuple[str, str, str]] = []  # (to, subject, html)

    async def send(self, to: str, subject: str, html: str) -> None:
        self.calls.append((to, subject, html))


@pytest.fixture
def fake_email_sender() -> FakeEmailSender:
    return FakeEmailSender()


@pytest.fixture
def fake_vector_store() -> FakeVectorStore:
    return FakeVectorStore()


@pytest.fixture
def make_chunk(db_session, user):
    from app.models import Chunk, Document

    counter = {"n": 0}

    def _make(content: str, owner=None) -> Chunk:
        owner = owner or user
        counter["n"] += 1
        n = counter["n"]
        doc = Document(
            user_id=owner.user_id,
            file_path=f"doc{n}.md",
            file_name=f"doc{n}.md",
            file_type="MD",
            hash_checksum=hashlib.sha256(f"doc{n}".encode()).hexdigest(),
        )
        db_session.add(doc)
        db_session.flush()
        chunk = Chunk(
            doc_id=doc.doc_id,
            chunk_index=0,
            chunk_hash=hashlib.sha256(f"chunk{n}".encode()).hexdigest(),
            content=content,
            token_count=len(content.split()),
        )
        db_session.add(chunk)
        db_session.commit()
        return chunk

    return _make
