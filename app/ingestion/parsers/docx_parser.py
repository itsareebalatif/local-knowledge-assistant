"""DOCX parser (python-docx).

python-docx exposes `document.paragraphs` and `document.tables` as two
separate flat lists, which loses their real interleaving order in the file.
We instead walk `document.element.body` directly and re-wrap each child as a
Paragraph or Table, so a table that sits between two paragraphs stays there
instead of being moved to the end.
"""

from __future__ import annotations

import io
import re

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.cleaning import normalize_text
from app.ingestion.exceptions import ParsingError
from app.ingestion.parsers.base import BaseParser
from app.ingestion.types import Block, ParsedDocument

_HEADING_STYLE = re.compile(r"^Heading\s*(\d)$", re.IGNORECASE)
_TITLE_STYLE = re.compile(r"^Title$", re.IGNORECASE)


def _iter_block_items(document: docx.Document):
    """Yield each top-level Paragraph/Table in document body order."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


class DocxParser(BaseParser):
    file_type = "DOCX"

    def parse_sync(self, file_bytes: bytes, filename: str) -> ParsedDocument:
        try:
            document = docx.Document(io.BytesIO(file_bytes))
        except Exception as exc:
            raise ParsingError(f"Could not open DOCX {filename}: {exc}") from exc

        blocks: list[Block] = []
        for item in _iter_block_items(document):
            if isinstance(item, Paragraph):
                text = normalize_text(item.text)
                if not text:
                    continue
                style_name = item.style.name if item.style else ""
                heading_match = _HEADING_STYLE.match(style_name or "")
                if heading_match or _TITLE_STYLE.match(style_name or ""):
                    level = int(heading_match.group(1)) if heading_match else 1
                    blocks.append(Block(type="heading", text=text, level=min(level, 6)))
                else:
                    blocks.append(Block(type="paragraph", text=text))
            else:  # Table
                rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in item.rows]
                text = normalize_text("\n".join(r for r in rows if r.strip()))
                if text:
                    blocks.append(Block(type="table", text=text))

        core_props = document.core_properties
        meta = {"source_format": "docx"}
        if core_props.title:
            meta["title"] = core_props.title
        if core_props.author:
            meta["author"] = core_props.author

        return ParsedDocument(blocks=blocks, meta=meta)
